"""验证模板驱动的工作簿入口：列序来自模板、迁移与回写语义。"""

from __future__ import annotations

import importlib.util
import tempfile
from pathlib import Path

import openpyxl

SCRIPT = Path(__file__).resolve().parents[1] / "scripts/literature-workbook-fallback.py"
spec = importlib.util.spec_from_file_location("workbook_fallback", SCRIPT)
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)

TEMPLATE = Path(__file__).resolve().parents[1] / "reference/文献检索汇总模板.xlsx"


def template_headers(path: Path = TEMPLATE) -> list[str]:
    return module.load_template(path).headers


def make_template(path: Path, headers: list[str]) -> Path:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "文献清单"
    sheet["A1"] = "文献检索汇总"
    sheet["A2"] = "测试模板"
    for column, header in enumerate(headers, start=1):
        sheet.cell(4, column, header)
    book.save(path)
    return path


def payload(**overrides) -> dict:
    record = {
        "rank": 1,
        "doi": "10.1038/s41524-021-00545-1",
        "title": "论文甲",
        "year": 2021,
        "journal": "期刊",
        "impact_factor": "2024 JCR 13.1",
        "cas_quartile": "2025版 1区",
        "journal_metrics_source": "测试来源",
    }
    record.update(overrides)
    return {"schema_version": 1, "topic": "测试", "records": [record]}


def test_reference_template_has_doi_and_template_order():
    headers = template_headers()
    assert headers[0] == "序号"
    assert "DOI" in headers
    assert headers.index("DOI") == 3
    assert "下载成功" in headers and "SI是否下载成功" in headers


def test_upsert_uses_template_columns():
    headers = template_headers()
    rows: list[dict[str, str]] = []
    resolved = module.upsert(rows, payload(), Path.cwd(), headers)

    assert len(rows) == 1
    row = rows[0]
    assert row["序号"] == 1
    assert row["DOI"] == "10.1038/s41524-021-00545-1"
    assert row["下载成功"] == "未尝试"
    assert row["SI是否下载成功"] == "未尝试"
    # 模板没有的列不应被写入
    assert "下载来源" not in row
    assert resolved[0]["sequence"] == 1
    assert resolved[0]["folder_name"].startswith("1_IF13.1_")


def test_upsert_keeps_existing_number_and_download_status():
    headers = template_headers()
    rows: list[dict[str, str]] = []
    module.upsert(rows, payload(), Path.cwd(), headers)
    folder = rows[0]["目录名"]
    rows[0]["下载成功"] = "是"

    module.upsert(rows, payload(title="更新题名"), Path.cwd(), headers)

    assert len(rows) == 1
    assert rows[0]["序号"] == 1
    assert rows[0]["目录名"] == folder
    assert rows[0]["下载成功"] == "是"
    assert rows[0]["论文题名"] == "更新题名"


def test_sequence_conflict_is_rejected():
    headers = template_headers()
    rows: list[dict[str, str]] = []
    module.upsert(rows, payload(), Path.cwd(), headers)

    try:
        module.upsert(rows, payload(rank=2, doi="10.1000/other", sequence=1), Path.cwd(), headers)
    except ValueError:
        return
    raise AssertionError("序号冲突应被拒绝")


def test_si_only_preserves_article_fields():
    headers = template_headers()
    rows: list[dict[str, str]] = []
    module.upsert(rows, payload(), Path.cwd(), headers)
    rows[0].update({"下载成功": "是", "下载状态": "cached", "PDF路径": "existing.pdf"})

    module.apply_report(
        rows,
        {
            "schema_version": 1,
            "results": [
                {
                    "doi": rows[0]["DOI"],
                    "download_mode": "supplements_only",
                    "status": "downloaded",
                    "supplement_status": "downloaded",
                    "success": True,
                }
            ],
        },
        headers,
    )

    assert rows[0]["SI是否下载成功"] == "是"
    assert rows[0]["PDF路径"] == "existing.pdf"
    assert rows[0]["下载状态"] == "cached"


def test_si_failure_keeps_saved_article_success():
    headers = template_headers()
    rows: list[dict[str, str]] = []
    module.upsert(rows, payload(), Path.cwd(), headers)
    with tempfile.TemporaryDirectory() as temporary:
        pdf = Path(temporary) / "article.pdf"
        pdf.write_bytes(b"%PDF-1.7\n" + b"x" * 2048)
        module.apply_report(
            rows,
            {
                "schema_version": 1,
                "results": [
                    {
                        "doi": rows[0]["DOI"],
                        "download_mode": "article_and_supplements",
                        "success": False,
                        "status": "failed",
                        "pdf_path": str(pdf),
                        "supplement_status": "unconfirmed",
                    }
                ],
            },
            headers,
        )

    assert rows[0]["下载成功"] == "是"
    assert rows[0]["SI是否下载成功"] == "否"


def test_custom_template_columns_drive_output(tmp_path: Path):
    custom = make_template(tmp_path / "t.xlsx", ["序号", "目录名", "论文题名", "DOI", "下载状态"])
    headers = template_headers(custom)
    rows: list[dict[str, str]] = []

    module.upsert(rows, payload(), Path.cwd(), headers)

    assert set(rows[0]) == set(headers)
    assert rows[0]["下载状态"] == "pending"


def test_migration_remaps_old_headers_into_template(tmp_path: Path):
    old_headers = ["序号", "目录名", "论文题名", "年份", "期刊", "下载成功", "下载状态"]
    old = make_template(tmp_path / "old.xlsx", old_headers)
    book = openpyxl.load_workbook(old)
    sheet = book["文献清单"]
    sheet.append([7, "7_旧目录", "旧题名", 2019, "旧期刊", "是", "downloaded"])
    book.save(old)

    book, sheet, headers, rows, migrated = module.load(old, TEMPLATE)

    assert migrated is True
    assert headers == template_headers()
    assert len(rows) == 1
    assert rows[0]["序号"] == 7
    assert rows[0]["论文题名"] == "旧题名"
    assert rows[0]["下载成功"] == "是"
    assert rows[0]["DOI"] == ""


def test_missing_required_column_reports_clearly(tmp_path: Path):
    custom = make_template(tmp_path / "t.xlsx", ["序号", "论文题名"])
    headers = template_headers(custom)

    try:
        module.upsert([], payload(), Path.cwd(), headers)
    except ValueError as error:
        assert "目录名" in str(error)
        return
    raise AssertionError("缺少必需列时应明确报错")


def test_save_writes_rows_into_template_layout(tmp_path: Path):
    headers = template_headers()
    rows: list[dict[str, str]] = []
    module.upsert(rows, payload(), Path.cwd(), headers)
    book, sheet, current_headers, _, _ = module.load(tmp_path / "汇总.xlsx", TEMPLATE)

    module.save(book, sheet, rows, current_headers, tmp_path / "汇总.xlsx")

    saved = openpyxl.load_workbook(tmp_path / "汇总.xlsx")
    data = saved["文献清单"]
    assert data.cell(4, 1).value == "序号"
    assert data.cell(4, 4).value == "DOI"
    assert data.cell(5, 1).value == 1
    assert data.cell(5, 4).value == "10.1038/s41524-021-00545-1"
