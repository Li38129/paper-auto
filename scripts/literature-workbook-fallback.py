"""按参考模板读写文献检索汇总工作簿（缺 Artifact Tool 时的入口）。

版式约定：`reference/文献检索汇总模板.xlsx` 是唯一版式来源，用户可以随时更新模板。
本脚本在运行时读取模板表头，按「内部字段 → 表头名称」的别名表写入，不硬编码列序；
模板缺少必需列时报错并说明缺哪一列。

命令：
  python scripts/literature-workbook-fallback.py --workbook 汇总.xlsx \
      --records records.json --resolved resolved.json
  python scripts/literature-workbook-fallback.py --workbook 汇总.xlsx --report batch-report.json
  python scripts/literature-workbook-fallback.py --workbook 汇总.xlsx --migrate-only true
"""

from __future__ import annotations

import argparse
import json
import os
import runpy
import sys
import tempfile
from datetime import UTC, datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

PROJECT = Path(__file__).resolve().parents[1]
# 单独加载现有纯函数，避免工作簿入口初始化网络下载依赖。
normalize_doi = runpy.run_path(str(PROJECT / "src/doi_harvester/doi.py"))["normalize_doi"]
literature_name = runpy.run_path(str(PROJECT / "src/doi_harvester/naming.py"))["literature_name"]
_SCHEMA = runpy.run_path(str(PROJECT / "src/doi_harvester/workbook_schema.py"))
build_column_index = _SCHEMA["build_column_index"]
detect_header_row = _SCHEMA["detect_header_row"]
load_template = _SCHEMA["load_template"]
missing_fields = _SCHEMA["missing_fields"]
read_headers = _SCHEMA["read_headers"]
HEADER_ALIASES = _SCHEMA["HEADER_ALIASES"]
REQUIRED_RECORD_FIELDS = _SCHEMA["REQUIRED_RECORD_FIELDS"]
REQUIRED_REPORT_FIELDS = _SCHEMA["REQUIRED_REPORT_FIELDS"]

# 交换记录字段 → 内部字段
RECORD_FIELDS = {
    "title": "title",
    "doi": "doi",
    "year": "year",
    "journal": "journal",
    "paper_type": "paper_type",
    "system": "system",
    "metric": "metric",
    "source_url": "source_url",
    "cas_quartile": "cas_quartile",
    "impact_factor": "impact_factor",
    "journal_metrics_source": "metrics_source",
}


def read_json(path: Path) -> dict:
    payload = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    if payload.get("schema_version") != 1:
        raise ValueError("交换数据必须使用 schema_version=1")
    return payload


def key_doi(value: object) -> str:
    return normalize_doi(str(value).strip()) if value else ""


def field_names(headers: list[str]) -> dict[str, str]:
    """内部字段 → 实际表头文本。"""
    index = build_column_index(headers)
    return {field: headers[position] for field, position in index.items()}


def read_rows(sheet: object, header_row: int, headers: list[str]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for values in sheet.iter_rows(
        min_row=header_row + 1, max_col=len(headers), values_only=True
    ):
        if any(value is not None and str(value).strip() for value in values):
            rows.append(
                {
                    header: ("" if value is None else value)
                    for header, value in zip(headers, values, strict=False)
                }
            )
    return rows


def remap_rows(
    rows: list[dict[str, str]], old_headers: list[str], new_headers: list[str]
) -> list[dict[str, str]]:
    """旧布局按表头名称搬到模板布局，同名保留，缺失置空。"""
    old_index = build_column_index(old_headers)
    new_index = build_column_index(new_headers)
    old_names = {field: old_headers[pos] for field, pos in old_index.items()}
    new_names = {field: new_headers[pos] for field, pos in new_index.items()}
    mapped: list[dict[str, str]] = []
    for row in rows:
        target = {header: "" for header in new_headers}
        for field, old_header in old_names.items():
            new_header = new_names.get(field)
            if new_header is not None:
                target[new_header] = row.get(old_header, "")
        mapped.append(target)
    return mapped


def validate_rows(rows: list[dict[str, str]], headers: list[str]) -> None:
    """序号与目录名必须唯一，DOI 不得重复。"""
    names = field_names(headers)
    sequences, folders, dois = set(), set(), set()
    for row in rows:
        sequence_value = row.get(names.get("rank", ""), "")
        try:
            sequence = int(sequence_value)
        except (TypeError, ValueError):
            raise ValueError(f"序号无效：{sequence_value}") from None
        folder = str(row.get(names.get("folder_name", ""), "")).casefold()
        doi = key_doi(row.get(names.get("doi", "")))
        if sequence < 1 or sequence in sequences or folder in folders:
            raise ValueError("已有记录的序号或目录冲突，原文件未修改")
        sequences.add(sequence)
        folders.add(folder)
        if doi:
            if doi in dois:
                raise ValueError("已有记录存在重复 DOI，原文件未修改")
            dois.add(doi)


def upsert(
    rows: list[dict[str, str]],
    payload: dict,
    folder_root: Path,
    headers: list[str],
) -> list[dict[str, object]]:
    names = field_names(headers)
    missing = missing_fields(headers, REQUIRED_RECORD_FIELDS)
    if missing:
        raise ValueError(f"模板缺少写入记录所需的列：{'、'.join(missing)}")
    validate_rows(rows, headers)

    records = payload.get("records", [])
    if not records:
        raise ValueError("records不能为空")
    ranks, seen_dois, seen_titles = set(), set(), set()
    for raw in records:
        rank, doi = raw.get("rank"), key_doi(raw.get("doi"))
        title = str(raw.get("title", "")).strip().casefold()
        if (
            not isinstance(rank, int)
            or isinstance(rank, bool)
            or rank < 1
            or rank in ranks
            or (doi and doi in seen_dois)
            or (not doi and title in seen_titles)
        ):
            raise ValueError("输入记录的rank、DOI或题名重复或无效")
        if not title and not doi:
            raise ValueError("输入记录缺少题名和DOI")
        ranks.add(rank)
        seen_dois.add(doi)
        seen_titles.add(title)

    stamp = datetime.now(UTC).isoformat()
    rank_header = names["rank"]
    max_seq = max(
        (int(row[rank_header]) for row in rows if str(row.get(rank_header, "")).isdigit()),
        default=0,
    )
    resolved: list[dict[str, object]] = []
    for raw in sorted(records, key=lambda item: item["rank"]):
        doi = key_doi(raw.get("doi"))
        title = str(raw.get("title", "")).strip()
        match: dict[str, str] | None = None
        for row in rows:
            row_doi = key_doi(row.get(names.get("doi", ""), "")) if "doi" in names else ""
            if doi and row_doi == doi:
                match = row
                break
            if (
                not doi
                and "title" in names
                and str(row.get(names["title"], "")).strip().casefold() == title.casefold()
            ):
                match = row
                break
        if match is not None:
            row = match
            existing_sequence = row.get(rank_header, "")
            if raw.get("sequence") is not None and int(raw["sequence"]) != int(existing_sequence):
                raise ValueError("导入记录与原编号不一致")
        else:
            sequence = raw.get("sequence", max_seq + 1)
            if (
                not isinstance(sequence, int)
                or isinstance(sequence, bool)
                or sequence < 1
                or any(str(row.get(rank_header, "")).strip() == str(sequence) for row in rows)
            ):
                raise ValueError("导入序号冲突或无效")
            folder = raw.get("folder_name") or literature_name(
                sequence,
                title,
                raw.get("year"),
                raw.get("impact_factor", ""),
                raw.get("cas_quartile", ""),
                raw.get("journal_abbreviation", ""),
            )
            if (
                Path(folder).name != folder
                or folder in {".", ".."}
                or any(
                    str(row.get(names["folder_name"], "")).casefold() == folder.casefold()
                    for row in rows
                )
            ):
                raise ValueError("目录名不安全或冲突")
            row = {header: "" for header in headers}
            row[rank_header] = sequence
            row[names["folder_name"]] = folder
            if "first_seen" in names:
                row[names["first_seen"]] = stamp
            if "download_success" in names:
                row[names["download_success"]] = "未尝试" if doi else "不适用"
            if "si_success" in names:
                row[names["si_success"]] = "未尝试" if doi else "不适用"
            if "download_status" in names:
                row[names["download_status"]] = "pending" if doi else "missing_doi"
            rows.append(row)
            max_seq = max(max_seq, sequence)

        for record_key, field in RECORD_FIELDS.items():
            value = raw.get(record_key)
            if value in (None, "") or field not in names:
                continue
            row[names[field]] = key_doi(value) if record_key == "doi" else value

        if "topic" in names:
            existing = str(row.get(names["topic"], "") or "")
            merged = list(
                dict.fromkeys(
                    part
                    for part in (
                        existing + "；" + str(payload.get("topic", ""))
                    ).split("；")
                    if part.strip()
                )
            )
            row[names["topic"]] = "；".join(merged)
        if "last_updated" in names:
            row[names["last_updated"]] = stamp

        resolved.append(
            {
                "rank": raw["rank"],
                "sequence": int(row[rank_header]),
                "folder_name": row.get(names["folder_name"], ""),
                "folder_path": str(
                    (folder_root / str(row.get(names["folder_name"], ""))).resolve()
                ),
                "title": row.get(names.get("title", ""), ""),
                "doi": row.get(names.get("doi", ""), ""),
            }
        )
    validate_rows(rows, headers)
    return resolved


def apply_report(rows: list[dict[str, str]], payload: dict, headers: list[str]) -> None:
    names = field_names(headers)
    missing = missing_fields(headers, REQUIRED_REPORT_FIELDS)
    if missing:
        raise ValueError(f"模板缺少回写报告所需的列：{'、'.join(missing)}")
    if "doi" not in names:
        raise ValueError("模板缺少「DOI」列，无法按 DOI 回写下载报告")
    mapping: dict[str, dict[str, str]] = {}
    for row in rows:
        doi = key_doi(row.get(names.get("doi", ""), ""))
        if doi:
            mapping[doi] = row
    results = payload.get("results", [])
    for result in results:
        doi = key_doi(result.get("doi"))
        row = mapping.get(doi)
        if row is None:
            raise ValueError(f"报告包含未汇总的DOI：{doi or '<空>'}")
        mode = result.get("download_mode", "article")
        if mode != "article" and "si_success" in names:
            supplement = str(result.get("supplement_status") or "not_requested")
            value = (
                "是"
                if supplement in {"downloaded", "cached"}
                else "未尝试"
                if supplement in {"not_requested", "pending"}
                else "否"
            )
            if result.get("status") == "auth_skipped":
                value = "否"
            row[names["si_success"]] = value
        if mode == "supplements_only":
            if "last_updated" in names:
                row[names["last_updated"]] = datetime.now(UTC).isoformat()
            continue
        status = str(result.get("status") or "failed")
        success = result.get("success") is True and status in {"downloaded", "cached"}
        validation = result.get("body_validation")
        if isinstance(validation, dict):
            success = validation.get("valid") is True
        elif result.get("pdf_path"):
            path = Path(str(result["pdf_path"]))
            try:
                with path.open("rb") as stream:
                    success = path.stat().st_size >= 1024 and stream.read(5) == b"%PDF-"
            except OSError:
                success = False
        if success and status not in {"downloaded", "cached"}:
            status = "downloaded"
        row[names["download_success"]] = "是" if success else "否"
        row[names["download_status"]] = status
        failure = "" if success else str(result.get("failure_reason") or status)
        row[names["failure_reason"]] = failure
        row[names["pdf_path"]] = result.get("pdf_path") or ""
        if "download_source" in names:
            row[names["download_source"]] = result.get("source") or ""
        if "last_updated" in names:
            row[names["last_updated"]] = datetime.now(UTC).isoformat()


def save(
    book: object,
    sheet: object,
    rows: list[dict[str, str]],
    headers: list[str],
    destination: Path,
) -> None:
    if not rows:
        raise ValueError("不导出空文献清单")
    header_row = detect_header_row(sheet)
    existing_last = sheet.max_row
    for offset, row in enumerate(rows):
        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(header_row + 1 + offset, column)
            cell.value = row.get(header, "")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    # 收缩时清掉遗留的旧行，避免旧数据残留在模板范围外
    for stale in range(header_row + 1 + len(rows), existing_last + 1):
        for column in range(1, len(headers) + 1):
            sheet.cell(stale, column).value = None
    if sheet.tables:
        for table in sheet.tables.values():
            table.ref = (
                f"A{header_row}:{get_column_letter(len(headers))}{header_row + len(rows)}"
            )
    destination.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(
        prefix=destination.stem + "-", suffix=".tmp.xlsx", dir=destination.parent
    )
    os.close(handle)
    try:
        book.save(temporary)
        os.replace(temporary, destination)
    finally:
        Path(temporary).unlink(missing_ok=True)


def load(workbook_path: Path, template_path: Path):
    """返回 (book, sheet, headers, rows, migrated)。"""
    template = load_template(template_path)
    if workbook_path.exists():
        book = openpyxl.load_workbook(workbook_path)
        sheet_name = (
            template.sheet_name
            if template.sheet_name in book.sheetnames
            else book.sheetnames[0]
        )
        sheet = book[sheet_name]
        header_row = detect_header_row(sheet)
        headers = read_headers(sheet, header_row)
        rows = read_rows(sheet, header_row, headers)
        if headers != template.headers:
            # 表头与模板不一致：按表头名称迁移到模板布局，保留可映射的数据
            target_book = openpyxl.load_workbook(template_path)
            target_sheet = target_book[template.sheet_name]
            return (
                target_book,
                target_sheet,
                template.headers,
                remap_rows(rows, headers, template.headers),
                True,
            )
        return book, sheet, headers, rows, False
    book = openpyxl.load_workbook(template_path)
    sheet = book[template.sheet_name]
    return book, sheet, template.headers, [], False


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument(
        "--template", type=Path, default=PROJECT / "reference/文献检索汇总模板.xlsx"
    )
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--records", type=Path)
    source.add_argument("--report", type=Path)
    source.add_argument("--migrate-only", nargs="?", const="true", default=None)
    parser.add_argument("--resolved", type=Path)
    args = parser.parse_args()

    if not args.template.is_file():
        print(f"参考模板不存在：{args.template}")
        return 1
    book, sheet, headers, rows, migrated = load(args.workbook, args.template)
    resolved: list[dict[str, object]] = []
    if args.records:
        if not args.resolved:
            parser.error("写入清单时必须指定--resolved")
        resolved = upsert(rows, read_json(args.records), args.workbook.parent.resolve(), headers)
    elif args.report:
        if migrated:
            print("提示：工作簿已按模板表头迁移，随后回写报告。")
        apply_report(rows, read_json(args.report), headers)
    save(book, sheet, rows, headers, args.workbook)
    if args.resolved and resolved:
        args.resolved.parent.mkdir(parents=True, exist_ok=True)
        args.resolved.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "workbook": str(args.workbook.resolve()),
                    "records": resolved,
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
    print(
        json.dumps(
            {
                "workbook": str(args.workbook.resolve()),
                "records": len(rows),
                "resolved": len(resolved),
                "migrated": migrated,
                "columns": len(headers),
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
