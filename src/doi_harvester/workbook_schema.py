"""以参考模板为准的工作簿表头解析与字段映射。

项目规则：`reference/文献检索汇总模板.xlsx` 是汇总工作簿的唯一版式来源，用户可以
随时更新模板。因此列的数量、顺序与名称都在运行时从模板读取，代码只按“内部字段 →
表头名称”的别名表写入，不硬编码列序。

模板缺少操作必需的列时，调用方应立即报错并说明缺哪一列，不得静默改版式。
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl

# 内部字段 → 可接受的表头写法（第一个为推荐写法）
HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "rank": ("序号",),
    "folder_name": ("目录名",),
    "title": ("论文题名", "题名"),
    "doi": ("DOI", "doi", "Doi"),
    "year": ("年份", "发表年份"),
    "journal": ("期刊", "期刊名称"),
    "paper_type": ("文献类型",),
    "system": ("体系/最佳组成", "体系或最佳组成"),
    "metric": ("关注指标",),
    "topic": ("检索主题",),
    "source_url": ("文献入口", "入口"),
    "download_success": ("下载成功",),
    "si_success": ("SI是否下载成功",),
    "download_status": ("下载状态",),
    "failure_reason": ("失败原因",),
    "pdf_path": ("PDF路径",),
    "download_source": ("下载来源",),
    "first_seen": ("首次收录时间",),
    "last_updated": ("最后更新时间",),
    "cas_quartile": ("中科院分区",),
    "impact_factor": ("影响因子（IF）", "影响因子(IF)", "影响因子"),
    "metrics_source": ("分区/IF依据", "分区或IF依据"),
}

# 写入记录所需的列
REQUIRED_RECORD_FIELDS = ("rank", "folder_name", "title")
# 回写下载报告所需的列
REQUIRED_REPORT_FIELDS = ("download_success", "download_status", "failure_reason", "pdf_path")


def normalize_header(value: object) -> str:
    """比较表头时忽略空白与大小写差异。"""
    return " ".join(str(value or "").replace("\u3000", " ").split()).strip().casefold()


def build_column_index(headers: list[str]) -> dict[str, int]:
    """把内部字段映射到 0 基列号；同一字段取第一个命中的别名。"""
    normalized = [normalize_header(item) for item in headers]
    index: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            key = normalize_header(alias)
            if key in normalized:
                index[field] = normalized.index(key)
                break
    return index


def missing_fields(headers: list[str], fields: tuple[str, ...]) -> list[str]:
    """返回缺失字段的推荐表头名称，用于明确报错。"""
    index = build_column_index(headers)
    return [HEADER_ALIASES[field][0] for field in fields if field not in index]


@dataclass(frozen=True, slots=True)
class TemplateSchema:
    """模板的表结构快照。"""

    path: Path
    sheet_name: str
    header_row: int
    headers: list[str]

    @property
    def columns(self) -> dict[str, int]:
        return build_column_index(self.headers)


def detect_header_row(sheet: object, *, scan_rows: int = 10, scan_columns: int = 40) -> int:
    """在模板或已有工作簿中定位表头行（含「序号」的首行）。"""
    limit = min(scan_columns, max(int(getattr(sheet, "max_column", 1) or 1), 1))
    for row in range(1, scan_rows + 1):
        values = [sheet.cell(row, column).value for column in range(1, limit + 1)]
        if any(normalize_header(value) == "序号" for value in values):
            return row
    raise ValueError("未找到表头行：前 10 行中没有包含「序号」的行。")


def read_headers(sheet: object, header_row: int) -> list[str]:
    """读取表头并去掉尾部空列。"""
    headers: list[str] = []
    for column in range(1, sheet.max_column + 1):
        value = sheet.cell(header_row, column).value
        headers.append(str(value).strip() if value is not None else "")
    while headers and not headers[-1]:
        headers.pop()
    if not headers:
        raise ValueError("表头为空，无法确定工作簿列。")
    return headers


def load_template(path: Path) -> TemplateSchema:
    """读取参考模板的表结构。"""
    workbook = openpyxl.load_workbook(path)
    sheet_name = workbook.sheetnames[0]
    sheet = workbook[sheet_name]
    header_row = detect_header_row(sheet)
    headers = read_headers(sheet, header_row)
    if "序号" not in headers:
        raise ValueError(f"模板缺少「序号」列：{path}")
    return TemplateSchema(path=path, sheet_name=sheet_name, header_row=header_row, headers=headers)
